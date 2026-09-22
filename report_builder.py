import os
import re
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# MARKDOWN PARSER
# ============================================================

def parse_report(markdown):

    lines = markdown.splitlines()

    sections = []
    current = None

    for line in lines:

        line = line.rstrip()

        # Main title
        if line.startswith("# ") and not line.startswith("## "):

            sections.append({
                "type": "title",
                "text": line[2:].strip(),
            })

            continue

        # Heading
        if line.startswith("## "):

            current = {
                "type": "section",
                "title": line[3:].strip(),
                "content": [],
            }

            sections.append(current)

            continue

        # Subheading
        if line.startswith("### "):

            if current is None:
                current = {
                    "type": "section",
                    "title": "",
                    "content": [],
                }
                sections.append(current)

            current["content"].append({
                "type": "subheading",
                "text": line[4:].strip(),
            })

            continue

        # Blank
        if not line.strip():

            if current:
                current["content"].append({
                    "type": "blank"
                })

            continue

        # Table
        if "|" in line:

            if current:

                current["content"].append({
                    "type": "raw",
                    "text": line,
                })

            continue

        # Normal text
        if current:

            current["content"].append({
                "type": "paragraph",
                "text": line,
            })

    return sections


# ============================================================
# EXTRACT TABLES
# ============================================================

def extract_tables(markdown):

    lines = markdown.splitlines()

    tables = []

    i = 0

    while i < len(lines):

        line = lines[i].strip()

        if (
            "|" in line
            and i + 1 < len(lines)
            and "|" in lines[i + 1]
            and "---" in lines[i + 1]
        ):

            headers = [
                x.strip()
                for x in line.strip("|").split("|")
            ]

            rows = []

            i += 2

            while i < len(lines):

                row = lines[i].strip()

                if not row or "|" not in row:
                    break

                values = [
                    x.strip()
                    for x in row.strip("|").split("|")
                ]

                rows.append(values)

                i += 1

            tables.append({
                "headers": headers,
                "rows": rows,
            })

            continue

        i += 1

    return tables


# ============================================================
# WORD HELPERS
# ============================================================

def set_cell_shading(cell, fill="1F4E78"):

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def set_cell_text_color(cell, color="FFFFFF"):

    for paragraph in cell.paragraphs:

        for run in paragraph.runs:

            run.font.color.rgb = __import__(
                "docx"
            ).shared.RGBColor.from_string(color)


def add_page_number(paragraph):

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_toc(document):

    paragraph = document.add_paragraph()

    run = paragraph.add_run()

    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_char_end)


def set_document_styles(document):

    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    for style_name, size in [
        ("Title", 28),
        ("Heading 1", 18),
        ("Heading 2", 14),
        ("Heading 3", 11.5),
    ]:

        style = styles[style_name]

        style.font.name = "Aptos"
        style.font.size = Pt(size)

        if style_name != "Title":
            style.font.bold = True


# ============================================================
# DOCX GENERATOR
# ============================================================

def create_docx(report, topic):

    document = Document()

    set_document_styles(document)

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # --------------------------------------------------------
    # COVER
    # --------------------------------------------------------

    for _ in range(4):
        document.add_paragraph()

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(topic)

    run.bold = True
    run.font.size = Pt(28)

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        "Professional Deep Research Report"
    )

    run.font.size = Pt(16)

    document.add_paragraph()

    date = document.add_paragraph()

    date.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date.add_run("Prepared by Deep Research AI Agent").bold = True

    date = document.add_paragraph()

    date.alignment = WD_ALIGN_PARAGRAPH.CENTER

    from datetime import datetime

    date.add_run(
        datetime.now().strftime("%d %B %Y")
    )

    document.add_page_break()

    # --------------------------------------------------------
    # DOCUMENT CONTROL
    # --------------------------------------------------------

    document.add_heading(
        "Document Control",
        level=1,
    )

    control = document.add_table(
        rows=4,
        cols=2,
    )

    control.alignment = WD_TABLE_ALIGNMENT.CENTER

    values = [
        ("Report", topic),
        ("Version", "1.0"),
        ("Status", "Final"),
        ("Generated", datetime.now().strftime("%d %B %Y")),
    ]

    for row, values_row in zip(control.rows, values):

        row.cells[0].text = values_row[0]
        row.cells[1].text = values_row[1]

        set_cell_shading(row.cells[0])

        set_cell_text_color(row.cells[0])

    document.add_page_break()

    # --------------------------------------------------------
    # TABLE OF CONTENTS
    # --------------------------------------------------------

    document.add_heading(
        "Table of Contents",
        level=1,
    )

    add_toc(document)

    document.add_page_break()

    # --------------------------------------------------------
    # REPORT
    # --------------------------------------------------------

    parsed = parse_report(report)

    table_counter = 0

    for item in parsed:

        if item["type"] == "title":
            continue

        if item["type"] == "section":

            title = item["title"]

            document.add_heading(
                title,
                level=1,
            )

            # Process section contents
            for content in item["content"]:

                ctype = content["type"]

                if ctype == "subheading":

                    document.add_heading(
                        content["text"],
                        level=2,
                    )

                elif ctype == "paragraph":

                    paragraph = document.add_paragraph(
                        content["text"]
                    )

                    paragraph.paragraph_format.space_after = Pt(6)
                    paragraph.paragraph_format.line_spacing = 1.15

                elif ctype == "raw":

                    # A table row will be processed later.
                    pass

    # --------------------------------------------------------
    # REAL TABLES
    # --------------------------------------------------------

    tables = extract_tables(report)

    if tables:

        document.add_page_break()

        document.add_heading(
            "Tables",
            level=1,
        )

        for table_data in tables:

            table_counter += 1

            caption = document.add_paragraph()

            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = caption.add_run(
                f"Table {table_counter}: Research Data"
            )

            run.bold = True

            data = [
                table_data["headers"]
            ] + table_data["rows"]

            table = document.add_table(
                rows=len(data),
                cols=len(data[0]),
            )

            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            table.style = "Table Grid"

            for r, row_data in enumerate(data):

                for c, value in enumerate(row_data):

                    cell = table.cell(r, c)

                    cell.text = str(value)

                    cell.vertical_alignment = (
                        WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    )

                    if r == 0:

                        set_cell_shading(cell)

                        set_cell_text_color(cell)

                        for run in cell.paragraphs[0].runs:
                            run.bold = True

            document.add_paragraph(
                "Source: As identified in the research report."
            )

            document.add_paragraph()

    # --------------------------------------------------------
    # HEADER / FOOTER
    # --------------------------------------------------------

    for section in document.sections:

        header = section.header

        p = header.paragraphs[0]

        p.text = topic

        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p.runs[0].font.size = Pt(8)

        footer = section.footer

        p = footer.paragraphs[0]

        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        r = p.add_run(
            "Deep Research AI Agent  |  Page "
        )

        r.font.size = Pt(8)

        add_page_number(p)

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# PDF
# ============================================================

def create_pdf(report, topic):

    output = BytesIO()

    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=20 * mm,
        bottomMargin=18 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        spaceAfter=20,
    )

    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=styles["Heading1"],
        fontSize=15,
        leading=19,
        spaceBefore=14,
        spaceAfter=8,
    )

    subheading_style = ParagraphStyle(
        "ReportSubheading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        spaceBefore=10,
        spaceAfter=5,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontSize=9.5,
        leading=14,
        spaceAfter=7,
    )

    story = []

    story.append(
        Spacer(1, 50 * mm)
    )

    story.append(
        Paragraph(
            topic,
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Professional Deep Research Report",
            styles["Heading2"],
        )
    )

    story.append(
        Spacer(1, 15 * mm)
    )

    from datetime import datetime

    story.append(
        Paragraph(
            datetime.now().strftime("%d %B %Y"),
            body_style,
        )
    )

    story.append(PageBreak())

    parsed = parse_report(report)

    for item in parsed:

        if item["type"] == "title":
            continue

        if item["type"] == "section":

            story.append(
                Paragraph(
                    item["title"],
                    heading_style,
                )
            )

            for content in item["content"]:

                if content["type"] == "subheading":

                    story.append(
                        Paragraph(
                            content["text"],
                            subheading_style,
                        )
                    )

                elif content["type"] == "paragraph":

                    text = content["text"]

                    text = (
                        text.replace("&", "&amp;")
                        .replace("<", "&lt;")
                        .replace(">", "&gt;")
                    )

                    story.append(
                        Paragraph(
                            text,
                            body_style,
                        )
                    )

    doc.build(story)

    output.seek(0)

    return output.getvalue()


# ============================================================
# XLSX
# ============================================================

def create_xlsx(report, topic):

    wb = Workbook()

    ws = wb.active

    ws.title = "Report Summary"

    ws["A1"] = topic

    ws["A1"].font = Font(
        bold=True,
        size=16,
    )

    ws["A3"] = "Generated Report"

    ws["A3"].font = Font(
        bold=True,
    )

    ws["A4"] = (
        "The workbook contains tables extracted "
        "from the research report."
    )

    tables = extract_tables(report)

    for index, table_data in enumerate(tables, start=1):

        sheet_name = f"Table_{index}"

        ws_table = wb.create_sheet(
            title=sheet_name[:31]
        )

        headers = table_data["headers"]

        for col, value in enumerate(headers, start=1):

            cell = ws_table.cell(
                row=1,
                column=col,
                value=value,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            cell.fill = __import__(
                "openpyxl"
            ).styles.PatternFill(
                start_color="1F4E78",
                end_color="1F4E78",
                fill_type="solid",
            )

        for row_index, row_data in enumerate(
            table_data["rows"],
            start=2,
        ):

            for col_index, value in enumerate(
                row_data,
                start=1,
            ):

                ws_table.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

        # Borders
        thin = Side(
            style="thin",
            color="D9E2F3",
        )

        for row in ws_table.iter_rows():

            for cell in row:

                cell.border = Border(
                    left=thin,
                    right=thin,
                    top=thin,
                    bottom=thin,
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        # Column widths
        for column in ws_table.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                if cell.value:

                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            ws_table.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                35,
            )

        ws_table.freeze_panes = "A2"

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()
